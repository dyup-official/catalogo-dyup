import openpyxl,json,re,sys,glob
from copy import copy
from openpyxl.worksheet.datavalidation import DataValidation
from hc import HC
src,dst=sys.argv[1],sys.argv[2]
h=open('catalogo.html').read()
itens=json.loads(re.search(r'const ITENS = (.*?);\n',h).group(1))
def faixa(p):
  return "Até R$ 800" if p<=800 else "R$ 800 a 1.200" if p<=1200 else "R$ 1.200 a 1.600" if p<=1600 else "R$ 1.600 a 2.000" if p<=2000 else "Acima de R$ 2.000"
brl=lambda v:f"{v:,.2f}".replace(",","X").replace(".",",").replace("X",".")
# agrupa por código HC + modelo + preço
grupos={}
for i in itens:
  for k in i['hc']:
    g=grupos.setdefault((k,i['marca'],i['modelo'],i['preco']),{"tam":set(),"cores":[]})
    g['tam'].update(i['tam']); g['cores'].append(f"{i['cor'] or 'cor única'} ({', '.join(i['refs'])})")
wb=openpyxl.load_workbook(src); ws=wb.worksheets[0]
# lê os blocos de categoria do modelo (linhas 5 a 34)
blocos=[];
for r in range(5,ws.max_row+1):
  code=ws.cell(r,1).value
  if not code: continue
  if not blocos or blocos[-1]['code']!=code:
    blocos.append({"code":code,"first":[copy(ws.cell(r,c)) for c in range(1,17)],"vals":[ws.cell(r,c).value for c in range(1,17)],
                   "cont":[ws.cell(r+1,c) for c in range(1,17)]})
# guarda estilos antes de apagar
for b in blocos:
  b['fs']=[ (c._style) for c in b['first']]
  b['cs']=[ copy(c._style) for c in b['cont']]
alt=ws.row_dimensions[5].height
ws.delete_rows(5,ws.max_row-4)
r=5
for b in blocos:
  linhas=sorted([(k,v) for k,v in grupos.items() if k[0]==b['code']],key=lambda kv:(kv[0][1],kv[0][2],kv[0][3]))
  n=max(len(linhas),1)
  for j in range(n):
    st=b['fs'] if j==0 else b['cs']
    for c in range(1,17): ws.cell(r,c)._style=copy(st[c-1])
    ws.cell(r,1).value=b['code']
    if j==0:
      for c in range(2,8): ws.cell(r,c).value=b['vals'][c-1]
    if linhas:
      (code,marca,modelo,preco),g=linhas[j]
      tams=sorted(g['tam'],key=lambda t:float(t.replace(',','.')))
      obs="Cores: "+"; ".join(g['cores'])
      extra=HC.get(modelo,([],""))[1]
      if extra: obs=extra+". "+obs
      vals=[marca.title() if marca!="ASICS" and marca!="HOKA" else marca, modelo, preco, faixa(preco), None, None, ", ".join(tams), None, obs]
      for c,v in zip(range(8,17),vals): ws.cell(r,c).value=v
      ws.cell(r,10).number_format='#,##0.00'
    else:
      ws.cell(r,16).value="Sem modelo desta categoria em estoque em "+sys.argv[3]
    r+=1
ws.data_validations.dataValidation=[]
dv=DataValidation(type="list",formula1='"Até R$ 800,R$ 800 a 1.200,R$ 1.200 a 1.600,R$ 1.600 a 2.000,Acima de R$ 2.000"',allow_blank=True)
dv.add(f"K5:K{r-1}"); ws.add_data_validation(dv)
wb.save(dst); print("linhas",r-5)
