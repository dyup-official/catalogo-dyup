# Planilha de revisão do catálogo (REVISAO DE PRODUTOS.xlsx, na pasta do catálogo no computador da loja).
#   python3 revisao.py ler <planilha.xlsx>                 -> grava excluir-revisao.txt com as refs marcadas em "Excluir?"
#   python3 revisao.py atualizar <entrada.xlsx> <saida.xlsx> <index.html> [fotos_novas.json]
#       mantém as marcações da pessoa, atualiza tamanhos/preço/situação da foto e acrescenta os produtos novos do estoque.
import sys,os,re,json,datetime,openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
AQUI=os.path.dirname(os.path.abspath(__file__))
CAB=['Categoria','Marca','Produto','Cor','Tamanhos','Preço','Referências','Excluir? (marque X)','Situação da foto','Trocar foto? (marque X)','Link ou observação da nova foto']
marcado=lambda v: v is not None and str(v).strip()!=''
refs_de=lambda v: [x.strip() for x in str(v or '').split(',') if x.strip()]
def linhas(f):
  ws=openpyxl.load_workbook(f)['No site']; return [list(r)+[None]*(len(CAB)-len(r)) for r in ws.iter_rows(min_row=2,values_only=True) if any(r)]
def ler(f):
  refs=[r for l in linhas(f) if marcado(l[7]) for r in refs_de(l[6])]
  open(os.path.join(AQUI,'excluir-revisao.txt'),'w').write('# gerado da planilha de revisão (coluna Excluir?)\n'+'\n'.join(refs)+'\n')
  print(len(refs),'refs excluídas pela planilha')
def atualizar(fe,fs,idx,novas=None):
  s=open(idx).read(); L=json.loads(re.search(r'ITENS\s*=\s*(\[.*?\]);\n',s,re.S).group(1))
  por_ref={}
  for i in L:
    for r in i['refs']: por_ref.setdefault(r,[]).append(i)
  pub=set(json.load(open(novas))) if novas and os.path.exists(novas) else set()  # refs com foto nova publicada agora
  hoje=datetime.date.today().strftime('%d/%m')
  velhas=linhas(fe) if fe and os.path.exists(fe) else []
  vistos=set(); out=[]
  for l in velhas:
    rs=refs_de(l[6])
    cand=[i for r in rs for i in por_ref.get(r,[]) if id(i) not in vistos]  # a mesma ref pode ter várias cores: confere o nome
    it=next((i for i in cand if i['modelo']==l[2] and (i['cor'] or None)==(l[3] or None)),None) or (cand[0] if len({id(i) for i in cand})==1 else None)
    if it:
      vistos.add(id(it)); l[4]=', '.join(it['tam']); l[5]=it['preco']
      for r in it['refs']:
        if r not in rs: rs.append(r)
      l[6]=', '.join(rs)
    if any(r in pub for r in rs): l[8]=f'Foto nova publicada {hoje}'; l[9]=None
    elif it: l[8]='Com foto' if it['foto'] else 'Sem foto'
    elif not marcado(l[7]): l[8]='Sem estoque'
    out.append(l[:len(CAB)])
  for it in L:
    if id(it) in vistos: continue
    out.append([it['grupo'],it['marca'],it['modelo'],it['cor'],', '.join(it['tam']),it['preco'],', '.join(it['refs']),None,
                'Com foto' if it['foto'] else 'Sem foto',None,None])
  wb=openpyxl.load_workbook(fe) if fe and os.path.exists(fe) else openpyxl.Workbook()
  if 'No site' in wb.sheetnames: idx_ws=wb.sheetnames.index('No site'); wb.remove(wb['No site'])
  else: idx_ws=0; [wb.remove(w) for w in list(wb.worksheets) if w.title=='Sheet']
  ws=wb.create_sheet('No site',idx_ws); ws.append(CAB)
  for l in out: ws.append(l)
  for c in ws[1]: c.font=Font(bold=True); c.fill=PatternFill('solid',fgColor='D0DF00'); c.alignment=Alignment(wrap_text=True,vertical='center')
  ws.freeze_panes='D2'; ws.auto_filter.ref=f"A1:K{ws.max_row}"
  for col,w in zip('ABCDEFGHIJK',[12,15,38,18,18,12,22,12,22,12,40]): ws.column_dimensions[col].width=w
  for c in ws['F'][1:]: c.number_format='"R$" #,##0.00'
  wb.active=idx_ws; wb.save(fs); print(len(out),'linhas;',len(out)-len(velhas),'novas')
if __name__=='__main__':
  if sys.argv[1]=='ler': ler(sys.argv[2])
  else: atualizar(*sys.argv[2:])
